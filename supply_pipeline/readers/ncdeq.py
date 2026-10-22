"""Method: spreadsheet download. NC DEQ publishes every construction stormwater certificate
of coverage as one xlsx. One GET, then a column mapping that has to survive format changes.

The mapping is header-aware: columns are found by their header text, never by position,
and the data sheet is the one whose first row has a "Permit Number" header. The publisher
has renamed the sheet and prepended columns before; a header that disappears raises with
the list of unmapped columns instead of shifting every value one column over.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
from pathlib import Path

from .. import http
from .base import ReaderResult

log = logging.getLogger(__name__)

METHOD = "spreadsheet"
EXAMPLE_FEED = 24   # Charlotte: NC DEQ NCG01 NOIs (xlsx)

URL = ("https://www.deq.nc.gov/energy-mineral-and-land-resources/stormwater/"
       "active-permits/activeexpired-construction-npdes-sw-cocs/download?attachment")

# table column -> header text before the first " (", lowercased
HEADER_MAP: dict[str, str] = {
    "permit_number": "permit number",
    "termination_date": "termination date",
    "originating_application": "originating application",
    "permit_type": "permit type",
    "general_permit": "general permit",
    "primary_permittee": "primary permittee",
    "project_name": "environmental interest",
    "full_address": "full address",
    "effective_date": "effective date",
    "expiration_date": "expiration date",
    "version": "version",
    "region": "region",
    "county": "county",
    "status": "status",
    "status_reason": "status reason",
    "created_on": "created on",
    "latitude": "latitude",
    "longitude": "longitude",
}
DATE_COLS = {"termination_date", "effective_date", "expiration_date"}
EXCEL_EPOCH = dt.datetime(1899, 12, 30)


class HeaderDrift(RuntimeError):
    """The sheet no longer has a header this reader maps."""


def norm_header(h) -> str:
    return "" if h is None else str(h).split(" (")[0].strip().lower()


def resolve_colmap(header_row: tuple) -> list[tuple[str, int]]:
    idx: dict[str, int] = {}
    for i, h in enumerate(header_row):
        n = norm_header(h)
        if n and n not in idx:
            idx[n] = i
    missing = [f"{col} ('{hdr}')" for col, hdr in HEADER_MAP.items() if hdr not in idx]
    if missing:
        raise HeaderDrift(f"unmapped columns: {missing}")
    return [(col, idx[hdr]) for col, hdr in HEADER_MAP.items()]


def find_data_sheet(wb):
    for ws in wb.worksheets:
        first = next(ws.iter_rows(values_only=True, max_row=1), None)
        if first and any(norm_header(h) == "permit number" for h in first):
            return ws
    raise HeaderDrift(f"no sheet with a 'Permit Number' header; sheets={wb.sheetnames}")


def excel_date(v, *, timestamp: bool = False) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, int | float):
        d = EXCEL_EPOCH + dt.timedelta(days=float(v))
        return d.isoformat() if timestamp else d.date().isoformat()
    if isinstance(v, dt.datetime):
        return v.isoformat() if timestamp else v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def to_record(row: tuple, colmap: list[tuple[str, int]], snapshot_date: str) -> dict | None:
    rec: dict = {}
    for col, i in colmap:
        v = row[i] if i < len(row) else None
        if col in DATE_COLS:
            v = excel_date(v)
        elif col == "created_on":
            v = excel_date(v, timestamp=True)
        elif col == "version":
            try:
                v = int(v) if v not in (None, "") else None
            except (TypeError, ValueError):
                v = None
        elif col in ("latitude", "longitude"):
            v = None if v in (None, "") else float(v)
        rec[col] = v
    if not rec.get("permit_number"):
        return None
    rec["snapshot_date"] = snapshot_date
    return rec


def parse(xlsx: bytes | Path, snapshot_date: str) -> tuple[list[dict], dict]:
    """Rows for NCG01 (construction stormwater) only; NCG25 is linear construction
    (pipelines, roads) and is not a housing supply signal."""
    import openpyxl
    src = io.BytesIO(xlsx) if isinstance(xlsx, bytes) else xlsx
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = find_data_sheet(wb)
    rows: list[dict] = []
    excluded: dict[str, int] = {}
    colmap = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            colmap = resolve_colmap(row)
            continue
        rec = to_record(row, colmap, snapshot_date)
        if rec is None:
            continue
        gp = (rec.get("general_permit") or "").strip()
        if not gp.upper().startswith("NCG01"):
            excluded[gp or "(blank)"] = excluded.get(gp or "(blank)", 0) + 1
            continue
        rows.append(rec)
    return rows, excluded


def download(session) -> bytes:
    r = session.get(URL, timeout=180)
    r.raise_for_status()
    if not r.content.startswith(b"PK"):
        raise RuntimeError(f"not an xlsx: {len(r.content)} bytes starting {r.content[:8]!r}")
    return r.content


def run(store, *, xlsx: bytes | Path | None = None, snapshot_date: str | None = None,
        session=None) -> ReaderResult:
    snapshot_date = snapshot_date or dt.date.today().isoformat()
    content = xlsx if xlsx is not None else download(session or http.session())
    rows, excluded = parse(content, snapshot_date)
    n = store.upsert("ncdeq.noi_submissions", rows, key=("permit_number",))
    return ReaderResult("ncdeq_noi_submissions", len(rows) + sum(excluded.values()), n,
                        notes={"excluded": excluded})
